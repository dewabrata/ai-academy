"""Sistem desain deck pelatihan: satu tempat yang menentukan rupa semua slide.

Dipisah dari `exporter.py` karena ini keputusan rupa, bukan keputusan isi.
`exporter.py` mengurai SLIDE.md dan memanggil `buat_deck()`; modul ini tidak
tahu apa-apa tentang Markdown.

Keputusan yang dipegang di sini, supaya tidak tersebar:

* **16:9.** Bawaan python-pptx adalah 10x7.5 inci (4:3) — format yang sudah
  pensiun sejak proyektor kelas berganti. Deck lama proyek ini semuanya 4:3.
* **Tidak memakai layout bawaan Office.** Semua slide dibangun dari layout
  kosong dan bentuk yang ditaruh sendiri. Placeholder bawaan membawa tema
  Office (Calibri abu-abu, tanpa warna) yang justru ingin kita tinggalkan.
* **Tanpa garis aksen di bawah judul dan tanpa bilah warna di tepi slide.**
  Keduanya penanda khas deck buatan mesin. Pemisahan dikerjakan oleh ruang
  kosong dan warna latar.
* **Satu motif yang diulang:** lingkaran aksen berisi nomor slide. Itu saja,
  dipakai konsisten di tiap slide isi.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# --- Palet ----------------------------------------------------------------
# Dipilih untuk pelatihan teknis: gelap arang yang mendominasi di slide sampul
# dan pembatas, putih bersih untuk isi, satu aksen hijau-laut. Sengaja bukan
# biru korporat, dan bukan krem.
TINTA = RGBColor(0x1B, 0x2A, 0x33)        # arang kebiruan - latar gelap
TINTA_MUDA = RGBColor(0x2C, 0x3E, 0x4A)
PUTIH = RGBColor(0xFF, 0xFF, 0xFF)
AKSEN = RGBColor(0x00, 0xA8, 0x96)        # hijau laut - motif & penekanan
AKSEN_TUA = RGBColor(0x02, 0x80, 0x90)
REDUP = RGBColor(0x62, 0x75, 0x7F)        # keterangan, nomor halaman
REDUP_TERANG = RGBColor(0x9E, 0xB0, 0xB9)
KODE_BG = RGBColor(0x14, 0x21, 0x2A)
KODE_FG = RGBColor(0xD6, 0xE4, 0xEA)

# --- Huruf ----------------------------------------------------------------
# Cambria dan Calibri ikut setiap pemasangan Office dan lebarnya terbaca sama
# di LibreOffice, jadi pemeriksaan luberan teks bisa dipercaya. Consolas untuk
# kode: ukurannya dihitung dari baris terpanjang, jadi tidak bergantung tebakan.
HURUF_JUDUL = "Cambria"
HURUF_ISI = "Calibri"
HURUF_KODE = "Consolas"

# --- Tata letak (inci) ----------------------------------------------------
LEBAR, TINGGI = 13.333, 7.5
TEPI = 0.8                                 # margin kiri/kanan
ATAS_JUDUL = 0.62
BADAN_ATAS = 1.95                          # tempat isi mulai di slide isi
BADAN_BAWAH = 6.75                         # batas bawah isi; sisanya nomor
JARAK = 0.35

MAKS_KODE_PT, MIN_KODE_PT = 13.0, 8.5


def _kotak(slide, x, y, lebar, tinggi):
    k = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(lebar), Inches(tinggi))
    tf = k.text_frame
    tf.word_wrap = True
    # Padding bawaan textbox menggeser teks beberapa milimeter dari posisi yang
    # dihitung; dinolkan supaya teks benar-benar lurus dengan bentuk lain.
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _tulis(tf, teks, *, ukuran, warna, huruf=HURUF_ISI, tebal=False,
           miring=False, spasi_setelah=0, rata=PP_ALIGN.LEFT, baru=False):
    p = tf.add_paragraph() if baru else tf.paragraphs[0]
    p.alignment = rata
    p.space_after = Pt(spasi_setelah)
    r = p.add_run()
    r.text = teks
    r.font.size = Pt(ukuran)
    r.font.bold = tebal
    r.font.italic = miring
    r.font.name = huruf
    r.font.color.rgb = warna
    return p


def _butir_bulat(p, warna=AKSEN):
    """Ganti bulatan bawaan dengan titik beraksen.

    python-pptx tidak membuka pengaturan bullet, jadi elemennya ditulis
    langsung. Tanpa ini bulatannya hitam dan ikut tema Office.
    """
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.28))))
    pPr.set("indent", str(-Emu(Inches(0.28))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for lama in pPr.findall(qn(tag)):
            pPr.remove(lama)
    warna_el = pPr.makeelement(qn("a:buClr"), {})
    srgb = warna_el.makeelement(qn("a:srgbClr"), {"val": str(warna)})
    warna_el.append(srgb)
    huruf_el = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    tanda = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    for el in (warna_el, huruf_el, tanda):
        pPr.append(el)


def _tanpa_bayangan(bentuk):
    """Matikan bayangan sungguhan, bukan hanya `shadow.inherit = False`.

    add_shape() menyertakan <p:style> yang menunjuk effectRef tema; selama
    rujukan itu ada, LibreOffice dan PowerPoint tetap menggambar bayangan
    meski daftar efeknya kosong. Lingkaran nomor dan kartu kode jadi tampak
    berhalo abu-abu - persis kesan "template bawaan" yang ingin dihindari.
    """
    bentuk.shadow.inherit = False
    gaya = bentuk._element.find(qn("p:style"))
    if gaya is not None:
        rujukan = gaya.find(qn("a:effectRef"))
        if rujukan is not None:
            rujukan.set("idx", "0")


def _latar(slide, warna):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = warna


def _lingkaran_nomor(slide, nomor, x, y, diameter=0.46, isi=AKSEN, teks=PUTIH):
    """Motif deck: satu lingkaran beraksen berisi nomor slide."""
    bentuk = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter))
    bentuk.fill.solid()
    bentuk.fill.fore_color.rgb = isi
    bentuk.line.fill.background()
    _tanpa_bayangan(bentuk)
    tf = bentuk.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _tulis(tf, str(nomor), ukuran=15, warna=teks, tebal=True, rata=PP_ALIGN.CENTER)
    return bentuk


def _baris_judul(judul, lebar_in, ukuran=31.0):
    """Berapa baris judul akan terpakai. Cambria rata-rata 0.5 em per huruf."""
    per_baris = max(int(lebar_in * 72.0 / (ukuran * 0.5)), 10)
    return max(1, -(-len(judul) // per_baris))


def _tinggi_butir(butir, ukuran, lebar_in):
    """Perkiraan tinggi daftar butir, dalam inci.

    Dipakai supaya kartu kode menempel di bawah butir alih-alih dilempar ke
    dasar ruang yang tersisa - yang meninggalkan lubang kosong di tengah
    slide. Lebar rata-rata huruf Calibri sekitar 0.48 em.
    """
    per_baris = max(int((lebar_in - 0.3) * 72.0 / (ukuran * 0.48)), 10)
    baris = sum(max(1, -(-len(b) // per_baris)) for b in butir)
    return (baris * ukuran * 1.22 + len(butir) * ukuran * 0.62) / 72.0


def _ukuran_kode(blok, lebar_in):
    """Ukuran huruf terbesar yang membuat baris terpanjang tetap muat.

    Lebar Consolas sekitar 0.55 em. Dihitung, bukan ditebak: satu baris yang
    meluber keluar kartu adalah cacat yang paling cepat terlihat di layar.
    """
    panjang = max((len(b) for b in blok), default=1) or 1
    pas = (lebar_in * 72.0) / (panjang * 0.55)
    return max(MIN_KODE_PT, min(MAKS_KODE_PT, pas))


def _kartu_kode(slide, blok, x, y, lebar, tinggi):
    kartu = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(lebar), Inches(tinggi))
    kartu.fill.solid()
    kartu.fill.fore_color.rgb = KODE_BG
    kartu.line.fill.background()
    _tanpa_bayangan(kartu)
    kartu.adjustments[0] = 0.04
    sisip = 0.22
    tf = _kotak(slide, x + sisip, y + sisip * 0.8,
                lebar - sisip * 2, tinggi - sisip * 1.6)
    ukuran = _ukuran_kode(blok, lebar - sisip * 2)
    for i, baris in enumerate(blok):
        _tulis(tf, baris or " ", ukuran=ukuran, warna=KODE_FG,
               huruf=HURUF_KODE, baru=i > 0)
    return kartu


def _tinggi_kode(blok, lebar_in):
    """Tinggi kartu yang cukup untuk semua barisnya, dalam inci."""
    ukuran = _ukuran_kode(blok, lebar_in - 0.44)
    return len(blok) * (ukuran * 1.25) / 72.0 + 0.5


# --- Jenis slide ----------------------------------------------------------
def _sampul(prs, judul, subjudul):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _latar(slide, TINTA)
    tf = _kotak(slide, TEPI, 2.35, LEBAR - TEPI * 2 - 1.2, 2.6)
    _tulis(tf, judul, ukuran=40, warna=PUTIH, huruf=HURUF_JUDUL, tebal=True,
           spasi_setelah=14)
    if subjudul:
        _tulis(tf, subjudul, ukuran=17, warna=REDUP_TERANG, baru=True)
    return slide


def _pembatas(prs, nomor, judul):
    """Slide tanpa butir dan tanpa kode diperlakukan sebagai pembatas bagian."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _latar(slide, TINTA)
    _lingkaran_nomor(slide, nomor, TEPI, 2.72, diameter=0.62)
    tf = _kotak(slide, TEPI + 0.95, 2.6, LEBAR - TEPI * 2 - 0.95, 1.9)
    _tulis(tf, judul, ukuran=32, warna=PUTIH, huruf=HURUF_JUDUL, tebal=True)
    return slide


def _nomor_halaman(slide, nomor):
    tf = _kotak(slide, LEBAR - TEPI - 1.0, TINGGI - 0.62, 1.0, 0.3)
    _tulis(tf, str(nomor), ukuran=10, warna=REDUP, rata=PP_ALIGN.RIGHT)


def _isi(prs, nomor, sl):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _latar(slide, PUTIH)
    _lingkaran_nomor(slide, nomor, TEPI, ATAS_JUDUL + 0.04)

    kiri = TEPI + 0.72
    lebar = LEBAR - kiri - TEPI
    judul = sl["judul"] or "(tanpa judul)"
    baris_judul = _baris_judul(judul, lebar)
    tinggi_judul = baris_judul * 31 * 1.22 / 72.0
    tf = _kotak(slide, kiri, ATAS_JUDUL, lebar, tinggi_judul + 0.1)
    _tulis(tf, judul, ukuran=31, warna=TINTA, huruf=HURUF_JUDUL, tebal=True)

    butir = sl["butir"]
    blok = [b for b in sl["kode"] if b]
    # Badan menempel di bawah judul. Titik mulai tetap membuat slide berjudul
    # satu baris menganga di atas, dan berjudul tiga baris berdesakan.
    y = max(BADAN_ATAS, ATAS_JUDUL + tinggi_judul + 0.42)
    tersedia = BADAN_BAWAH - y

    if blok:
        # Kode mengambil tinggi yang ia butuhkan, butir memakai sisanya. Deck
        # lama menaruh kartu kode di titik tetap 2/3 tinggi slide sehingga ia
        # menimpa daftar butir di hampir setiap slide yang punya keduanya.
        perlu = sum(_tinggi_kode(b, lebar) for b in blok) + JARAK * len(blok)
        perlu = min(perlu, tersedia - (1.1 if butir else 0))
    else:
        perlu = 0.0

    if butir:
        ukuran = 18 if len(butir) <= 4 else (16 if len(butir) <= 6 else 14)
        # Kalau tidak ada kode dan ruangnya longgar, hurufnya dinaikkan
        # daripada menyisakan separuh slide kosong di bawah.
        if not blok:
            while ukuran < 24 and _tinggi_butir(butir, ukuran + 1, lebar) < tersedia:
                ukuran += 1
        tinggi_butir = min(_tinggi_butir(butir, ukuran, lebar), tersedia - perlu)
        tfb = _kotak(slide, kiri, y, lebar, max(tinggi_butir, 0.6))
        for i, b in enumerate(butir):
            p = _tulis(tfb, b, ukuran=ukuran, warna=TINTA_MUDA,
                       spasi_setelah=ukuran * 0.62, baru=i > 0)
            _butir_bulat(p)
        y += max(tinggi_butir, 0.6) + JARAK

    for b in blok:
        tinggi = min(_tinggi_kode(b, lebar), max(BADAN_BAWAH - y, 0.6))
        if tinggi < 0.5:
            break
        _kartu_kode(slide, b, kiri, y, lebar, tinggi)
        y += tinggi + JARAK

    _nomor_halaman(slide, nomor)
    return slide


def buat_deck(judul: str, slides: list[dict], keluar: Path,
              subjudul: str = "Materi pelatihan") -> Path:
    """Tulis deck ber-desain ke `keluar`. `slides` dari exporter.parse_slides()."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(LEBAR), Inches(TINGGI)

    _sampul(prs, judul, subjudul)
    for i, sl in enumerate(slides, 1):
        kosong = not sl["butir"] and not any(sl["kode"])
        s = _pembatas(prs, i, sl["judul"] or "(tanpa judul)") if kosong else _isi(prs, i, sl)
        if sl["catatan"]:
            s.notes_slide.notes_text_frame.text = "\n".join(sl["catatan"])

    keluar.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(keluar))
    return keluar
