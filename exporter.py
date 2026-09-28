"""Ubah sumber Markdown menjadi DOCX dan PPTX.

Peran AI hanya menulis Markdown; berkas biner selalu turunan dari sumber itu.
Konsekuensi yang disengaja: revisi apa pun mengubah Markdown lalu biner
dihasilkan ulang, jadi DOCX/PPTX tidak pernah menyimpang dari sumbernya dan
hasil revisi tidak perlu ditelusuri di dua tempat.

Dipanggil academy.py setelah paket tiap pertemuan dibangun dan setelah revisi.
Bisa juga dijalankan sendiri:

    python exporter.py workspace/<proyek>
"""
import re
import sys
from pathlib import Path

# Batas kepadatan slide — nilainya juga ditulis di prompts/slide.md supaya peran
# Slide tahu targetnya. Kalau diubah di sini, ubah juga di sana.
MAKS_BUTIR = 6
MAKS_KATA_BUTIR = 14
MAKS_BARIS_KODE = 10


class ExportError(Exception):
    """Konversi gagal. Ini kegagalan tahap, bukan peringatan yang bisa diabaikan."""


# ---------------------------------------------------------------------------
# Markdown -> DOCX
# ---------------------------------------------------------------------------
_INLINE = [
    (re.compile(r"\*\*(.+?)\*\*"), r"\1"),
    (re.compile(r"(?<!\w)\*(.+?)\*(?!\w)"), r"\1"),
    (re.compile(r"`(.+?)`"), r"\1"),
    (re.compile(r"\[(.+?)\]\((.+?)\)"), r"\1 (\2)"),
]

# Satu pola untuk semua penanda inline sekaligus, supaya urutannya di dalam
# kalimat tetap terjaga saat dipecah menjadi run.
_POTONG = re.compile(
    r"\*\*(?P<tebal>.+?)\*\*"
    r"|(?<!\w)\*(?P<miring>.+?)\*(?!\w)"
    r"|`(?P<kode>.+?)`"
    r"|\[(?P<teks>.+?)\]\((?P<url>.+?)\)")

KODE_FONT = "Consolas"
KODE_LATAR = "F3F4F6"          # abu sangat muda, tetap terbaca saat dicetak
KODE_GARIS = "9CA3AF"


def _bersih(teks: str) -> str:
    """Buang penanda inline Markdown. Dipakai di tempat yang hanya menerima
    teks polos: judul, butir slide, dan catatan pengajar."""
    for pola, ganti in _INLINE:
        teks = pola.sub(ganti, teks)
    return teks


def _latar(el, warna: str):
    """Pasang <w:shd> pada elemen paragraf atau run."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), warna)
    el.append(shd)


def _gaya_kode(doc):
    """Style paragraf blok kode: font lebar tetap, latar abu, garis tepi kiri.

    Tanpa latar dan garis, blok kode di Word hanya berbeda fontnya — pembaca
    tidak melihatnya sebagai blok, dan itulah keluhan yang memicu perubahan ini.
    """
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, Inches

    gaya = doc.styles.add_style("KodeLab", 1)   # 1 = WD_STYLE_TYPE.PARAGRAPH
    gaya.font.name = KODE_FONT
    gaya.font.size = Pt(9)
    pf = gaya.paragraph_format
    pf.left_indent = Inches(0.22)
    pf.right_indent = Inches(0.1)
    pf.space_before = Pt(6)
    pf.space_after = Pt(8)
    pf.keep_together = True

    # Word mengabaikan font.name untuk sebagian skrip kalau rFonts tidak diisi.
    rpr = gaya.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for atribut in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(atribut), KODE_FONT)

    ppr = gaya.element.get_or_add_pPr()
    _latar(ppr, KODE_LATAR)
    tepi = OxmlElement("w:pBdr")
    kiri = OxmlElement("w:left")
    kiri.set(qn("w:val"), "single")
    kiri.set(qn("w:sz"), "18")
    kiri.set(qn("w:space"), "8")
    kiri.set(qn("w:color"), KODE_GARIS)
    tepi.append(kiri)
    ppr.append(tepi)
    return gaya


def _tulis_inline(par, teks: str):
    """Isi paragraf dengan run yang mempertahankan tebal, miring, dan kode.

    Inline code jadi run berfont lebar tetap. Sebelumnya backtick-nya dibuang
    tanpa pengganti, sehingga nama berkas dan perintah di tengah kalimat tidak
    bisa dibedakan dari kata biasa.
    """
    from docx.oxml.ns import qn
    from docx.shared import Pt

    posisi = 0
    for m in _POTONG.finditer(teks):
        if m.start() > posisi:
            par.add_run(teks[posisi:m.start()])
        if m.group("tebal") is not None:
            par.add_run(m.group("tebal")).bold = True
        elif m.group("miring") is not None:
            par.add_run(m.group("miring")).italic = True
        elif m.group("kode") is not None:
            r = par.add_run(m.group("kode"))
            r.font.name = KODE_FONT
            r.font.size = Pt(9.5)
            rfonts = r._element.get_or_add_rPr().get_or_add_rFonts()
            for atribut in ("w:ascii", "w:hAnsi", "w:cs"):
                rfonts.set(qn(atribut), KODE_FONT)
            _latar(r._element.get_or_add_rPr(), KODE_LATAR)
        else:
            par.add_run(f"{m.group('teks')} ({m.group('url')})")
        posisi = m.end()
    if posisi < len(teks):
        par.add_run(teks[posisi:])
    return par


def md_ke_docx(md: Path, keluar: Path) -> Path:
    try:
        from docx import Document
    except ImportError as e:
        raise ExportError("Butuh python-docx: pip install -r requirements.txt") from e

    if not md.exists():
        raise ExportError(f"Sumber tidak ada: {md}")
    isi = md.read_text(encoding="utf-8")
    if not isi.strip():
        raise ExportError(f"Sumber kosong: {md}")

    doc = Document()
    _gaya_kode(doc)

    dalam_kode = False
    baris_kode: list[str] = []
    ada_isi = False
    # Markdown membungkus paragraf di ~80 kolom. Tanpa penyangga ini tiap baris
    # sumber menjadi paragraf Word tersendiri, sehingga satu paragraf utuh
    # tampil terpotong-potong dengan jarak di antaranya.
    teks_buf: list[str] = []
    gaya_buf: str | None = None

    def tutup_kode():
        nonlocal baris_kode
        if baris_kode:
            doc.add_paragraph("\n".join(baris_kode), style="KodeLab")
            baris_kode = []

    def tutup_teks():
        nonlocal teks_buf, gaya_buf
        if teks_buf:
            par = doc.add_paragraph(style=gaya_buf) if gaya_buf else doc.add_paragraph()
            _tulis_inline(par, " ".join(teks_buf))
            teks_buf = []
        gaya_buf = None

    for baris in isi.splitlines():
        if baris.lstrip().startswith("```"):
            if dalam_kode:
                tutup_kode()
            else:
                tutup_teks()
            dalam_kode = not dalam_kode
            continue
        if dalam_kode:
            baris_kode.append(baris)
            ada_isi = True
            continue

        b = baris.rstrip()
        if not b.strip():
            tutup_teks()
            continue

        m = re.match(r"^(#{1,6})\s+(.*)", b)
        if m:
            tutup_teks()
            doc.add_heading(_bersih(m.group(2)), level=min(len(m.group(1)), 4))
            ada_isi = True
            continue
        # Baris tabel dipertahankan apa adanya sebagai satu paragraf: mengubahnya
        # jadi tabel Word akan memaksa menebak lebar kolom, dan tabel di modul
        # umumnya hanya dua-tiga kolom pendek.
        if b.lstrip().startswith("|"):
            tutup_teks()
            if re.fullmatch(r"[\s|:-]+", b):
                continue          # baris pemisah header tabel
            _tulis_inline(doc.add_paragraph(),
                          b.strip().strip("|").replace("|", " \u00b7 "))
            ada_isi = True
            continue
        m = re.match(r"^\s*[-*+]\s+(.*)", b)
        if m:
            tutup_teks()
            teks_buf, gaya_buf = [m.group(1)], "List Bullet"
            ada_isi = True
            continue
        m = re.match(r"^\s*(\d+)[.)]\s+(.*)", b)
        if m:
            tutup_teks()
            teks_buf, gaya_buf = [m.group(2)], "List Number"
            ada_isi = True
            continue
        if b.lstrip().startswith(">"):
            isi_kutipan = b.lstrip().lstrip(">").strip()
            if gaya_buf != "Intense Quote":
                tutup_teks()
                gaya_buf = "Intense Quote"
            if isi_kutipan:
                teks_buf.append(isi_kutipan)
            ada_isi = True
            continue
        # Baris biasa: lanjutan paragraf, butir, atau kutipan yang sedang
        # dibangun — atau awal paragraf baru kalau tidak ada yang terbuka.
        teks_buf.append(b.strip())
        ada_isi = True

    tutup_teks()
    tutup_kode()
    if not ada_isi:
        raise ExportError(f"Tidak ada isi yang bisa dikonversi dari {md.name}")

    keluar.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(keluar))
    except Exception as e:
        raise ExportError(f"Gagal menyimpan {keluar.name}: {e}") from e
    return keluar


# ---------------------------------------------------------------------------
# Markdown -> PPTX
# ---------------------------------------------------------------------------
def parse_slides(isi: str) -> tuple[str, list[dict]]:
    """Urai SLIDE.md menjadi (judul deck, daftar slide).

    Kontraknya ada di prompts/slide.md: `#` sekali di atas, `##` = satu slide,
    isinya butir / blok kode / baris `> catatan pengajar:`.
    """
    judul = ""
    slides: list[dict] = []
    kini: dict | None = None
    dalam_kode = False

    for baris in isi.splitlines():
        if baris.lstrip().startswith("```"):
            dalam_kode = not dalam_kode
            if kini is not None:
                if dalam_kode:
                    kini["kode"].append([])
                # Penutup blok tidak perlu diapa-apakan; barisnya sudah terkumpul.
            continue
        if dalam_kode:
            if kini is not None and kini["kode"]:
                kini["kode"][-1].append(baris)
            continue

        m = re.match(r"^##\s+(.*)", baris)
        if m:
            kini = {"judul": _bersih(m.group(1)).strip(), "butir": [], "kode": [], "catatan": []}
            slides.append(kini)
            continue
        m = re.match(r"^#\s+(.*)", baris)
        if m and not judul:
            judul = _bersih(m.group(1)).strip()
            continue
        if kini is None:
            continue

        b = baris.strip()
        if not b:
            continue
        m = re.match(r"^>\s*catatan pengajar:\s*(.*)", b, re.I)
        if m:
            kini["catatan"].append(m.group(1).strip())
            continue
        if b.startswith(">"):
            kini["catatan"].append(_bersih(b.lstrip(">").strip()))
            continue
        m = re.match(r"^[-*+]\s+(.*)", b)
        if m:
            kini["butir"].append(_bersih(m.group(1)).strip())
            continue
        # Paragraf lepas tidak sesuai kontrak, tapi membuangnya berarti isi hilang
        # tanpa jejak. Jadikan butir dan biarkan pemeriksa kepadatan menegurnya.
        kini["butir"].append(_bersih(b))

    return judul, slides


def periksa_kepadatan(slides: list[dict], sumber: Path) -> list[str]:
    """Temuan pelanggaran batas kepadatan, satu baris per pelanggaran."""
    temuan = []
    for i, s in enumerate(slides, 1):
        nama = f"slide {i} ('{s['judul'][:40]}')"
        if len(s["butir"]) > MAKS_BUTIR:
            temuan.append(f"{nama}: {len(s['butir'])} butir, batas {MAKS_BUTIR}")
        for b in s["butir"]:
            n = len(b.split())
            if n > MAKS_KATA_BUTIR:
                temuan.append(f"{nama}: butir {n} kata, batas {MAKS_KATA_BUTIR} — \"{b[:60]}...\"")
        for blok in s["kode"]:
            if len(blok) > MAKS_BARIS_KODE:
                temuan.append(f"{nama}: blok kode {len(blok)} baris, batas {MAKS_BARIS_KODE}")
    return temuan


def md_ke_pptx(md: Path, keluar: Path) -> tuple[Path, list[str]]:
    """Hasilkan PPTX. Mengembalikan (path, temuan kepadatan).

    Temuan kepadatan TIDAK menggagalkan konversi: deck tetap dibuat supaya bisa
    dilihat, dan temuannya naik ke daftar perbaikan lewat Reviewer. Yang
    menggagalkan adalah sumber yang tidak bisa diurai sama sekali.
    """
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except ImportError as e:
        raise ExportError("Butuh python-pptx: pip install -r requirements.txt") from e

    if not md.exists():
        raise ExportError(f"Sumber tidak ada: {md}")
    judul, slides = parse_slides(md.read_text(encoding="utf-8"))
    if not slides:
        raise ExportError(
            f"{md.name}: tidak ada slide. Setiap slide harus dimulai dengan '## '."
        )

    prs = Presentation()
    tata_judul = prs.slide_layouts[0]       # judul + subjudul
    tata_isi = prs.slide_layouts[1]         # judul + isi
    tata_kosong = prs.slide_layouts[5]      # judul saja

    s = prs.slides.add_slide(tata_judul)
    s.shapes.title.text = judul or md.parent.name
    if len(s.placeholders) > 1:
        s.placeholders[1].text = "Materi pelatihan"

    for sl in slides:
        punya_teks = bool(sl["butir"])
        slide = prs.slides.add_slide(tata_isi if punya_teks else tata_kosong)
        slide.shapes.title.text = sl["judul"] or "(tanpa judul)"

        if punya_teks:
            tf = slide.placeholders[1].text_frame
            tf.clear()
            for i, b in enumerate(sl["butir"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = b
                p.level = 0
                p.font.size = Pt(20 if len(sl["butir"]) <= 4 else 16)

        for blok in sl["kode"]:
            if not blok:
                continue
            kotak = slide.shapes.add_textbox(
                prs.slide_width // 12, prs.slide_height * 2 // 3,
                prs.slide_width * 5 // 6, prs.slide_height // 4)
            tf = kotak.text_frame
            tf.word_wrap = True
            for i, baris in enumerate(blok):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = baris
                p.font.name = "Consolas"
                p.font.size = Pt(12)

        if sl["catatan"]:
            slide.notes_slide.notes_text_frame.text = "\n".join(sl["catatan"])

    keluar.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(keluar))
    except Exception as e:
        raise ExportError(f"Gagal menyimpan {keluar.name}: {e}") from e
    return keluar, periksa_kepadatan(slides, md)


# ---------------------------------------------------------------------------
# Satu pertemuan / satu proyek
# ---------------------------------------------------------------------------
def gabung_handbook(folder: Path, judul: str) -> Path:
    """Bentuk HANDBOOK.md dari point/point-*.md, tanpa mengubah isi point.

    Digabung di Python, bukan oleh peran: handbook adalah gabungan point yang
    sudah lolos telaah, jadi menulis ulang isinya hanya membuka peluang merusak
    yang sudah benar — dan tidak butuh model sama sekali.
    """
    point = sorted((folder / "point").glob("point-[0-9]*.md"))
    point = [p for p in point if re.fullmatch(r"point-\d+\.md", p.name)]
    if not point:
        raise ExportError(f"Tidak ada point di {folder / 'point'}")
    isi = [p.read_text(encoding="utf-8").strip() for p in point]
    daftar = []
    for i, teks in enumerate(isi, 1):
        m = re.search(r"^##\s+(.+)$", teks, re.M)
        # Nama lokalnya TIDAK boleh `judul`: itu nama parameter judul pertemuan,
        # dan menimpanya membuat judul handbook menjadi judul point terakhir.
        judul_point = m.group(1).strip() if m else point[i - 1].stem
        # Judul point sudah memuat nomornya sendiri ("## 2. Setup ..."). Tanpa
        # ini daftar isi menulis nomor dua kali: "2. 2. Setup ...".
        bernomor = re.match(r"^(\d+(?:\.\d+)*)[.)]\s+", judul_point)
        daftar.append(f"{bernomor.group(1)}. {judul_point[bernomor.end():]}" if bernomor
                      else f"{i}. {judul_point}")
    keluar = folder / "HANDBOOK.md"
    keluar.write_text(
        f"# {judul}\n\n## Daftar isi\n\n" + "\n".join(daftar) + "\n\n"
        + "\n\n".join(isi) + "\n", encoding="utf-8")
    return keluar


def ekspor_pertemuan(folder: Path, docx: bool = True,
                     pptx: bool = True) -> tuple[list[Path], list[str]]:
    """Hasilkan ulang biner untuk satu folder pertemuan.

    `docx`/`pptx` salah = konversi dilewati; sumber Markdown-nya tetap ada dan
    tetap bisa dikonversi belakangan dengan `python exporter.py <ruang kerja>`.

    Mengembalikan (berkas yang dibuat, temuan kepadatan). Melempar ExportError
    kalau sumber yang ada gagal dikonversi.
    """
    dibuat: list[Path] = []
    temuan: list[str] = []

    for nama in ("HANDBOOK", "LATIHAN", "KUNCI", "PRAKTIK"):
        md = folder / f"{nama}.md"
        if docx and md.exists():
            dibuat.append(md_ke_docx(md, folder / f"{nama}.docx"))
    slide = folder / "SLIDE.md"
    if pptx and slide.exists():
        path, t = md_ke_pptx(slide, folder / "SLIDE.pptx")
        dibuat.append(path)
        temuan += [f"{folder.name}: {x}" for x in t]
    return dibuat, temuan


def ekspor_proyek(ws: Path) -> tuple[list[Path], list[str], list[str]]:
    """Hasilkan ulang biner untuk seluruh pertemuan.

    Mengembalikan (berkas dibuat, temuan kepadatan, kegagalan). Kegagalan
    dikumpulkan alih-alih dilempar, supaya satu pertemuan rusak tidak
    menyembunyikan hasil pertemuan lain.
    """
    dibuat: list[Path] = []
    temuan: list[str] = []
    gagal: list[str] = []
    for folder in sorted((ws / "materi").glob("pertemuan-*")):
        if not folder.is_dir():
            continue
        try:
            d, t = ekspor_pertemuan(folder)
            dibuat += d
            temuan += t
        except ExportError as e:
            gagal.append(f"{folder.name}: {e}")
    return dibuat, temuan, gagal


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    ws = Path(sys.argv[1])
    if not (ws / "materi").exists():
        print(f"Bukan ruang kerja proyek (tidak ada {ws / 'materi'}).")
        sys.exit(2)
    dibuat, temuan, gagal = ekspor_proyek(ws)
    for p in dibuat:
        print(f"  dibuat  {p}")
    for t in temuan:
        print(f"  padat   {t}")
    for g in gagal:
        print(f"  GAGAL   {g}")
    print(f"\n{len(dibuat)} berkas, {len(temuan)} temuan kepadatan, {len(gagal)} gagal.")
    sys.exit(1 if gagal else 0)
